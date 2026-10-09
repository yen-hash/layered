#!/usr/bin/env python3
"""Builds public/downloads/layered-renovation-budget-planner.xlsx (needs: pip install openpyxl).
Developer tool only: the generated file is committed, so production never runs this."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import os

GREEN, TINT, YELLOW, GREY = '0A7360', 'E4EFEA', 'FFF4CC', 'F3F5F1'
hdr_font = Font(bold=True, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor=GREEN)
inp_fill = PatternFill('solid', fgColor=YELLOW)
tot_fill = PatternFill('solid', fgColor=TINT)
thin = Side(style='thin', color='C9D6D0')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
money = '"S$"#,##0'
pct = '0%'

wb = Workbook()

# ---------------- Read me ----------------
rm = wb.active
rm.title = 'Read me'
lines = [
    ('Layered renovation budget planner', 'title'),
    ('layeredsg.com/resources', 'link'),
    ('', None),
    ('How to use this workbook', 'h'),
    ('1. On the Budget sheet, fill the yellow cells. Start with your total budget ceiling (the most you will spend), then split it across categories in the Planned column.', None),
    ('2. When quotes arrive, type each firm\'s price per category into Quote A, B and C, then enter the quote you chose in the Chosen column. The Difference column shows how it compares with your plan.', None),
    ('3. As work is paid for, record what you really spend in the Actual column. Update it after each stage, not at the end.', None),
    ('4. Keep a buffer (10% to 15% is common). It is added for you below the totals. Treat it as spent only when a change is agreed in writing.', None),
    ('5. List the things outside the renovation contract (furniture, appliances, curtains, moving, cleaning) in the second table so they do not quietly raid the renovation budget.', None),
    ('6. Use the Quote comparison sheet to line up two or three firms on the same scope.', None),
    ('', None),
    ('Good to know', 'h'),
    ('There are no prices in this workbook. Enter your own quotes. For indicative 2026 ranges, use the cost calculator and the itemised estimator on layeredsg.com.', None),
    ('The "typical share" column is only a rough guide to how renovation budgets often split. Your home will differ: a flat with little carpentry will shift the split.', None),
    ('GST: GST-registered firms charge GST (9% in 2026). Check whether each quote includes it, and set the GST cell on the Budget sheet to match. Check the current rate with IRAS.', None),
    ('This workbook is general information for planning, not financial or professional advice.', None),
]
for i, (text, kind) in enumerate(lines, start=1):
    c = rm.cell(row=i, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if kind == 'title': c.font = Font(bold=True, size=18, color=GREEN)
    elif kind == 'h': c.font = Font(bold=True, size=12, color=GREEN)
    elif kind == 'link': c.font = Font(color=GREEN, underline='single'); c.hyperlink = 'https://layeredsg.com/resources'
rm.column_dimensions['A'].width = 110

# ---------------- Budget ----------------
b = wb.create_sheet('Budget')
b['A1'] = 'Renovation budget'; b['A1'].font = Font(bold=True, size=16, color=GREEN)
b['A3'] = 'Total budget ceiling (S$)'; b['C3'] = None
b['A4'] = 'Buffer for changes and surprises'; b['C4'] = 0.10
b['A5'] = 'GST rate'; b['C5'] = 0.09
b['A6'] = 'Do the quotes already include GST?'; b['C6'] = 'No'
for r in (3, 4, 5, 6):
    b.cell(row=r, column=1).font = Font(bold=True)
    b.cell(row=r, column=3).fill = inp_fill; b.cell(row=r, column=3).border = box
b['C3'].number_format = money; b['C4'].number_format = pct; b['C5'].number_format = '0.0%'
b['D3'] = 'The most you will spend, including the buffer.'; b['D4'] = '10% to 15% is common; resale homes often need more.'
b['D5'] = 'Check the current rate with IRAS.'; b['D6'] = 'Choose Yes or No.'
for r in (3, 4, 5, 6): b.cell(row=r, column=4).font = Font(italic=True, color='5C6B65')
dv = DataValidation(type='list', formula1='"Yes,No"', allow_blank=False); b.add_data_validation(dv); dv.add('C6')

HEAD = 9
heads = ['Category', 'Typical share (rough guide)', 'Planned (S$)', 'Quote A (S$)', 'Quote B (S$)', 'Quote C (S$)', 'Chosen (S$)', 'Actual spent (S$)', 'Chosen minus planned', 'Notes']
for j, h in enumerate(heads, start=1):
    c = b.cell(row=HEAD, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = box
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
b.row_dimensions[HEAD].height = 34

cats = [
    ('Hacking and disposal', ''),
    ('Carpentry and built-ins (wardrobes, cabinets)', 'About a third'),
    ('Flooring and tiling', '10% to 15%'),
    ('Kitchen (wet works, counters, fittings)', '15% to 20% with bathrooms'),
    ('Bathrooms (wet works, fittings, screens)', '15% to 20% with kitchen'),
    ('Electrical and rewiring', '10% to 12% with plumbing'),
    ('Plumbing', '10% to 12% with electrical'),
    ('False ceiling and lighting', '10% to 15% with painting and doors'),
    ('Painting', '10% to 15% with ceilings and doors'),
    ('Doors, gates and windows', '10% to 15% with ceilings and painting'),
    ('Aircon installation', ''),
    ('Design, project management and permits', ''),
    ('Other renovation works', ''),
]
first = HEAD + 1
for i, (name, share) in enumerate(cats):
    r = first + i
    b.cell(row=r, column=1, value=name)
    b.cell(row=r, column=2, value=share).font = Font(color='5C6B65')
    for col in range(3, 9):
        c = b.cell(row=r, column=col); c.fill = inp_fill; c.number_format = money
    b.cell(row=r, column=9, value=f'=IF(OR(C{r}="",G{r}=""),"",G{r}-C{r})').number_format = '"S$"#,##0;[Red]-"S$"#,##0'
    b.cell(row=r, column=10).fill = inp_fill
    for col in range(1, 11): b.cell(row=r, column=col).border = box
last = first + len(cats) - 1
tot = last + 1
b.cell(row=tot, column=1, value='Renovation total').font = Font(bold=True)
for col in range(3, 9):
    L = get_column_letter(col)
    c = b.cell(row=tot, column=col, value=f'=SUM({L}{first}:{L}{last})'); c.number_format = money
b.cell(row=tot, column=9, value=f'=IF(C{tot}=0,"",G{tot}-C{tot})').number_format = '"S$"#,##0;[Red]-"S$"#,##0'
for col in range(1, 11):
    c = b.cell(row=tot, column=col); c.fill = tot_fill; c.font = Font(bold=True); c.border = box

# summary
s = tot + 2
rows = [
    ('Planned total plus buffer', f'=C{tot}*(1+C4)', money),
    ('GST on planned total (0 if already included in your figures)', f'=IF(C6="Yes",0,C{tot}*(1+C4)*C5)', money),
    ('Planned renovation total with buffer and GST', f'=C{s}+C{s+1}', money),
    ('Within your ceiling?', f'=IF(C3="","Enter your ceiling in C3",IF(C{s+2}<=C3,"Yes, within the ceiling","Over the ceiling by S$"&TEXT(C{s+2}-C3,"#,##0")))', None),
    ('Chosen quotes with GST (0 added if already included)', f'=IF(C6="Yes",G{tot},G{tot}*(1+C5))', money),
    ('Spent so far (actual)', f'=H{tot}', money),
]
for k, (label, formula, fmt) in enumerate(rows):
    r = s + k
    b.cell(row=r, column=1, value=label).font = Font(bold=True)
    c = b.cell(row=r, column=3, value=formula)
    if fmt: c.number_format = fmt
    c.font = Font(bold=True); c.fill = tot_fill
    for col in (1, 2, 3): b.cell(row=r, column=col).border = box

# outside the contract
o = s + len(rows) + 2
b.cell(row=o, column=1, value='Outside the renovation contract').font = Font(bold=True, size=13, color=GREEN)
oh = o + 1
for j, h in enumerate(['Item', '', 'Planned (S$)', 'Quote / price (S$)', '', '', '', 'Actual spent (S$)', 'Actual minus planned', 'Notes'], start=1):
    if h:
        c = b.cell(row=oh, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = box
        c.alignment = Alignment(wrap_text=True, horizontal='center')
outs = ['Furniture (sofa, dining table, beds)', 'Appliances (fridge, washer, hob and hood if not in contract)', 'Curtains and blinds', 'Lighting fittings bought separately', 'Moving', 'Post-renovation cleaning', 'Renovation insurance (if not in contract)', 'Other']
of = oh + 1
for i, name in enumerate(outs):
    r = of + i
    b.cell(row=r, column=1, value=name)
    for col in (3, 4, 8): b.cell(row=r, column=col).fill = inp_fill; b.cell(row=r, column=col).number_format = money
    b.cell(row=r, column=10).fill = inp_fill
    b.cell(row=r, column=9, value=f'=IF(OR(C{r}="",H{r}=""),"",H{r}-C{r})').number_format = '"S$"#,##0;[Red]-"S$"#,##0'
    for col in range(1, 11): b.cell(row=r, column=col).border = box
ol = of + len(outs) - 1
ot = ol + 1
b.cell(row=ot, column=1, value='Outside-contract total').font = Font(bold=True)
for col in (3, 4, 8):
    L = get_column_letter(col)
    c = b.cell(row=ot, column=col, value=f'=SUM({L}{of}:{L}{ol})'); c.number_format = money
for col in range(1, 11):
    c = b.cell(row=ot, column=col); c.fill = tot_fill; c.font = Font(bold=True); c.border = box
g = ot + 2
b.cell(row=g, column=1, value='Everything: planned renovation (with buffer and GST) plus outside items').font = Font(bold=True)
c = b.cell(row=g, column=3, value=f'=C{s+2}+C{ot}'); c.number_format = money; c.font = Font(bold=True); c.fill = tot_fill

widths = {1: 52, 2: 30, 3: 16, 4: 16, 5: 16, 6: 16, 7: 16, 8: 16, 9: 20, 10: 34}
for col, w in widths.items(): b.column_dimensions[get_column_letter(col)].width = w
b.freeze_panes = f'B{HEAD + 1}'
b.page_setup.orientation = 'landscape'; b.page_setup.fitToWidth = 1; b.page_setup.fitToHeight = 0
b.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------- Quote comparison ----------------
q = wb.create_sheet('Quote comparison')
q['A1'] = 'Compare quotes on the same scope'; q['A1'].font = Font(bold=True, size=16, color=GREEN)
q['A2'] = 'Give every firm the same written brief, then fill one column per firm. Compare the final totals, not the headline numbers.'
q['A2'].font = Font(italic=True, color='5C6B65')
for j, h in enumerate(['', 'Firm A', 'Firm B', 'Firm C'], start=1):
    c = q.cell(row=4, column=j, value=h); c.font = hdr_font; c.fill = hdr_fill; c.border = box; c.alignment = Alignment(horizontal='center')
qrows = [
    ('Company name', None), ('Contact person and phone', None), ('Quote date', None),
    ('Quoted total (S$)', 'money'), ('Does the quote include GST? (Yes / No)', None),
    ('Total including GST (S$)', 'calc'),
    ('Total vs your planned total', 'diff'),
    ('Itemised with quantities and materials? (Yes / No)', None),
    ('Carpentry foot runs and finish stated', None),
    ('Key materials and brands (tiles, laminates, hardware)', None),
    ('What is excluded', None),
    ('Payment schedule (stages and amounts)', None),
    ('Start date and expected completion', None),
    ('Delay clause', None),
    ('Warranty and how to claim', None),
    ('Who supervises the site, and how often', None),
    ('HDB licence / CaseTrust checked on the official lookups? (Yes / No)', None),
    ('Anything that worries you', None),
]
r0 = 5
tot_row = r0 + 3  # Quoted total
gst_row = r0 + 4
for i, (label, kind) in enumerate(qrows):
    r = r0 + i
    q.cell(row=r, column=1, value=label).font = Font(bold=True)
    q.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical='top')
    q.cell(row=r, column=1).border = box
    for col in (2, 3, 4):
        L = get_column_letter(col)
        c = q.cell(row=r, column=col); c.border = box; c.alignment = Alignment(wrap_text=True, vertical='top')
        if kind == 'money': c.fill = inp_fill; c.number_format = money
        elif kind == 'calc':
            c.value = f'=IF({L}{tot_row}="","",IF({L}{gst_row}="Yes",{L}{tot_row},{L}{tot_row}*(1+Budget!$C$5)))'; c.number_format = money; c.fill = tot_fill; c.font = Font(bold=True)
        elif kind == 'diff':
            c.value = f'=IF(OR({L}{r0+5}="",Budget!$C${tot}=0),"",{L}{r0+5}-Budget!$C${s+2})'; c.number_format = '"S$"#,##0;[Red]-"S$"#,##0'; c.fill = tot_fill
        else: c.fill = inp_fill
dv2 = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True); q.add_data_validation(dv2)
for i, (label, kind) in enumerate(qrows):
    if '(Yes / No)' in label:
        for col in (2, 3, 4): dv2.add(f'{get_column_letter(col)}{r0 + i}')
q.column_dimensions['A'].width = 52
for L in 'BCD': q.column_dimensions[L].width = 34
q.freeze_panes = 'B5'
q.page_setup.orientation = 'landscape'; q.page_setup.fitToWidth = 1; q.page_setup.fitToHeight = 0
q.sheet_properties.pageSetUpPr.fitToPage = True

out = os.path.join(os.path.dirname(__file__), '..', '..', 'public', 'downloads', 'layered-renovation-budget-planner.xlsx')
wb.save(os.path.normpath(out))
print('wrote', os.path.normpath(out), 'tot row', tot, 'summary row', s)
